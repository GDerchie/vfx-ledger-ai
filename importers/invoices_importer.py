"""
importers/invoices_importer.py — Validated invoice import.
"""
from importers.base import ImportResult, validate_cost, build_diff

INV_COL_MAP = {
    'vendor': 'vendor', 'vendor name': 'vendor',
    'ep': 'episode', 'episode': 'episode',
    'invoice': 'invoice_num', 'invoice #': 'invoice_num', 'invoice_num': 'invoice_num',
    'inv num': 'invoice_num', 'inv_num': 'invoice_num',
    'date': 'invoice_date', 'invoice date': 'invoice_date',
    'amount': 'amount', 'total': 'amount', 'gross': 'amount',
    'status': 'status',
    'notes': 'notes',
    'po': 'po_number', 'po number': 'po_number', 'po_number': 'po_number',
}
INV_DB_FIELDS = ['vendor', 'episode', 'invoice_num', 'invoice_date',
                 'amount', 'status', 'notes', 'po_number']
VALID_STATUSES = {'PAID', 'PENDING', 'APPROVED', 'DISPUTED', 'CANCELLED'}


def run(file_bytes: bytes, ep: int, conn, mode: str = 'upsert',
        dry_run: bool = False) -> ImportResult:
    result = ImportResult()
    try:
        import openpyxl, io
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active

        headers = []
        header_row = None
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if len([c for c in row if c is not None]) >= 3:
                headers = [str(c).strip().lower() if c else '' for c in row]
                header_row = i
                break
        if header_row is None:
            result.warnings.append("No header row found")
            return result

        col_to_field = {idx: INV_COL_MAP[h]
                        for idx, h in enumerate(headers) if h in INV_COL_MAP}

        rows_out = []
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if all(c is None for c in row):
                continue
            rec = {}
            for idx, f in col_to_field.items():
                val = row[idx] if idx < len(row) else None
                if val is not None:
                    rec[f] = str(val).strip() if not isinstance(val, (int, float)) else val
            if rec.get('vendor') or rec.get('invoice_num'):
                if 'episode' not in rec:
                    rec['episode'] = ep
                rows_out.append(rec)

        result.rows_parsed = len(rows_out)

        for row in rows_out:
            # Validate amount
            if 'amount' in row:
                v = validate_cost(row['amount'], 'amount', result.warnings)
                if v is None:
                    row.pop('amount')
                else:
                    row['amount'] = v

            # Normalise status
            if 'status' in row:
                s = str(row['status']).strip().upper()
                if s not in VALID_STATUSES:
                    result.warnings.append(
                        f"Invoice {row.get('invoice_num','?')}: unknown status '{s}' — set to PENDING"
                    )
                    row['status'] = 'PENDING'
                else:
                    row['status'] = s

            existing = None
            if row.get('invoice_num') and row.get('vendor'):
                ex = conn.execute(
                    'SELECT * FROM invoice_log WHERE UPPER(vendor)=UPPER(?) AND invoice_num=?',
                    (row['vendor'], row['invoice_num'])
                ).fetchone()
                existing = dict(ex) if ex else None

            diff = build_diff(existing, row, ['id', 'created_at'])
            if diff:
                diff['invoice_num'] = row.get('invoice_num', '')
                result.diff.append(diff)

            if dry_run:
                result.rows_inserted += (1 if diff and diff['action'] == 'insert' else 0)
                result.rows_updated  += (1 if diff and diff['action'] == 'update'  else 0)
                result.rows_skipped  += (1 if not diff else 0)
                continue

            try:
                if existing is None:
                    fields = [k for k in row if k in INV_DB_FIELDS]
                    ph = ', '.join('?' * len(fields))
                    conn.execute(
                        f"INSERT INTO invoice_log ({', '.join(fields)}) VALUES ({ph})",
                        [row[f] for f in fields]
                    )
                    result.rows_inserted += 1
                elif mode != 'append' and diff:
                    updates = {c['field']: c['new'] for c in diff['changes']
                               if c['field'] in INV_DB_FIELDS}
                    if updates:
                        sc = ', '.join(f'{k}=?' for k in updates)
                        conn.execute(
                            f'UPDATE invoice_log SET {sc} WHERE id=?',
                            list(updates.values()) + [existing['id']]
                        )
                        result.rows_updated += 1
                    else:
                        result.rows_skipped += 1
                else:
                    result.rows_skipped += 1
            except Exception as e:
                result.errors.append(f"Invoice {row.get('invoice_num','?')}: {e}")

        if not dry_run:
            conn.commit()

    except Exception as e:
        result.errors.append(f"Parse error: {e}")

    return result
