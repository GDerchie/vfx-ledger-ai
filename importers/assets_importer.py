"""
importers/assets_importer.py — Validated asset import.
"""
from importers.base import ImportResult, validate_cost, build_diff

ASSET_COL_MAP = {
    'asset': 'asset_name', 'asset name': 'asset_name', 'asset_name': 'asset_name',
    'type': 'asset_type', 'asset type': 'asset_type', 'asset_type': 'asset_type',
    'ep': 'ep', 'episode': 'ep',
    'sc': 'scene_code', 'scene': 'scene_code', 'scene_code': 'scene_code',
    'description': 'description', 'desc': 'description',
    'turnover': 'turnover', 'turn': 'turnover',
    'notes': 'notes',
    'repeat': 'repeat_asset', 'repeat asset': 'repeat_asset',
    'lidar': 'lidar',
    'photo': 'photogrammetry', 'photogrammetry': 'photogrammetry',
    'est': 'est_budget', 'est budget': 'est_budget', 'est_budget': 'est_budget',
    'actual': 'actual_spend', 'actual spend': 'actual_spend', 'actual_spend': 'actual_spend',
    'award': 'award_vendor', 'vendor': 'award_vendor', 'award_vendor': 'award_vendor',
}
ASSET_DB_FIELDS = ['ep', 'scene_code', 'asset_name', 'asset_type', 'description',
                   'turnover', 'notes', 'repeat_asset', 'lidar', 'photogrammetry',
                   'est_budget', 'actual_spend', 'award_vendor']


def run(file_bytes: bytes, ep: int, conn, mode: str = 'upsert',
        sheet_hint: str = None, dry_run: bool = False) -> ImportResult:
    result = ImportResult()
    try:
        import openpyxl, io
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        for sh in wb.sheetnames:
            if str(ep) in sh:
                ws = wb[sh]
                break

        header_row = None
        headers    = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if len([c for c in row if c is not None]) >= 3:
                headers = [str(c).strip().lower() if c else '' for c in row]
                header_row = i
                break
        if header_row is None:
            result.warnings.append("No header row found")
            return result

        col_to_field = {idx: ASSET_COL_MAP[h]
                        for idx, h in enumerate(headers) if h in ASSET_COL_MAP}

        rows_out = []
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if all(c is None for c in row):
                continue
            rec = {'ep': ep}
            for idx, f in col_to_field.items():
                val = row[idx] if idx < len(row) else None
                if val is not None:
                    rec[f] = str(val).strip() if not isinstance(val, (int, float)) else val
            if rec.get('asset_name') or rec.get('scene_code'):
                rows_out.append(rec)

        result.rows_parsed = len(rows_out)

        for row in rows_out:
            for cf in ('est_budget', 'actual_spend'):
                if cf in row and row[cf] is not None:
                    v = validate_cost(row[cf], cf, result.warnings)
                    if v is None:
                        row.pop(cf)
                    else:
                        row[cf] = v

            existing = None
            if row.get('asset_name') and row.get('scene_code'):
                ex = conn.execute(
                    'SELECT * FROM assets WHERE ep=? AND asset_name=? AND scene_code=?',
                    (ep, row['asset_name'], row['scene_code'])
                ).fetchone()
                existing = dict(ex) if ex else None

            diff = build_diff(existing, row, ['id', 'ep', 'created_at'])
            if diff:
                diff['asset_name'] = row.get('asset_name', '')
                result.diff.append(diff)

            if dry_run:
                if diff:
                    result.rows_inserted += (1 if diff['action'] == 'insert' else 0)
                    result.rows_updated  += (1 if diff['action'] == 'update' else 0)
                else:
                    result.rows_skipped += 1
                continue

            try:
                if existing is None:
                    fields = [k for k in row if k in ASSET_DB_FIELDS]
                    ph = ', '.join('?' * len(fields))
                    conn.execute(
                        f"INSERT INTO assets ({', '.join(fields)}) VALUES ({ph})",
                        [row[f] for f in fields]
                    )
                    result.rows_inserted += 1
                elif mode != 'append' and diff:
                    updates = {c['field']: c['new'] for c in diff['changes']
                               if c['field'] in ASSET_DB_FIELDS}
                    if updates:
                        sc = ', '.join(f'{k}=?' for k in updates)
                        conn.execute(
                            f'UPDATE assets SET {sc} WHERE id=?',
                            list(updates.values()) + [existing['id']]
                        )
                        result.rows_updated += 1
                    else:
                        result.rows_skipped += 1
                else:
                    result.rows_skipped += 1
            except Exception as e:
                result.errors.append(f"Asset {row.get('asset_name','?')}: {e}")

        if not dry_run:
            conn.commit()

    except Exception as e:
        result.errors.append(f"Parse error: {e}")

    return result
