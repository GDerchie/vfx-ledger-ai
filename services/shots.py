"""
services/shots.py — Shot data helpers, Excel import, column maps.
"""
import io
from services.db import get_db, _MAX_COST

COL_MAP = {
    'ep':'ep','id':'shot_num','shot_num':'shot_num','shot #':'shot_num','#':'shot_num',
    'sc':'scene_code','scene_code':'scene_code','scene #':'scene_code',
    's_code':'s_code','shot code':'s_code',
    'loc':'location','location':'location',
    'ext/int':'ext_int','ext_int':'ext_int','int/ext':'ext_int',
    'pg':'pg','page':'pg','img':'img','image':'img',
    'asset':'asset','assets':'asset',
    'script description':'script_desc','script_desc':'script_desc','script desc':'script_desc',
    'vfx description':'vfx_desc','vfx_desc':'vfx_desc','vfx desc':'vfx_desc',
    'type':'shot_type','shot_type':'shot_type',
    'complexity':'complexity','complex':'complexity',
    'omit':'omit',
    'overall shot history  showrunner/director/ep/pd notes \nfor internal purpose only':'history_notes',
    'overall shot history  showrunner/director/ep/pd notes \nfor internal purpose only ':'history_notes',
    'history':'history_notes','history_notes':'history_notes','shot history':'history_notes',
    'director notes':'history_notes',
    'shot(est)':'shot_est','shot est':'shot_est','shot_est':'shot_est','shots':'shot_est',
    'cost(est)':'cost_est','cost est':'cost_est','cost_est':'cost_est','cost':'cost_est',
    'budget(est)':'budget_est','budget est':'budget_est','budget_est':'budget_est','budget':'budget_est',
    'award(vendor)':'award_vendor','award vendor':'award_vendor','award_vendor':'award_vendor','vendor':'award_vendor',
    'edit_count':'edit_count','edit count':'edit_count','edit #':'edit_count','edit_#':'edit_count',
    'budget(awardcost)':'budget_award_cost','budget award cost':'budget_award_cost',
    'efc':'efc',
    'variance':None,
    'notes':'notes',
}

DB_FIELDS = ['shot_num','scene_code','s_code','location','ext_int','pg','img',
             'asset','script_desc','vfx_desc','shot_type','complexity','omit',
             'history_notes','shot_est','cost_est','budget_est','award_vendor',
             'edit_count','budget_award_cost','efc','notes']


def _parse_rows_from_excel(file_bytes, ep, sheet_hint=None):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    target = None
    if sheet_hint and sheet_hint in wb.sheetnames:
        target = wb[sheet_hint]
    if not target:
        ep_str = str(ep)
        for name in wb.sheetnames:
            if ep_str in name:
                target = wb[name]
                break
    if not target:
        target = wb.active

    header_row_idx = None
    headers = []
    for i, row in enumerate(target.iter_rows(values_only=True), 1):
        if len([c for c in row if c is not None]) >= 4:
            headers = [str(c).strip().lower() if c is not None else '' for c in row]
            header_row_idx = i
            break
    if header_row_idx is None:
        return [], wb.sheetnames

    col_to_field = {}
    for idx, h in enumerate(headers):
        field = COL_MAP.get(h)
        if field:
            col_to_field[idx] = field

    rows_out = []
    for row in target.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if all(c is None for c in row):
            continue
        rec = {'ep': int(ep)}
        for idx, field in col_to_field.items():
            if field == 'ep':
                continue
            val = row[idx] if idx < len(row) else None
            if val is not None:
                if field == 'omit':
                    rec[field] = 1 if str(val).strip().lower() in ('true','1','yes','x','omit') else 0
                elif field in ('shot_est', 'edit_count'):
                    try: rec[field] = int(float(val))
                    except: rec[field] = 0
                elif field in ('cost_est', 'budget_est', 'budget_award_cost', 'efc'):
                    try: rec[field] = float(val)
                    except: rec[field] = 0.0
                else:
                    rec[field] = str(val).strip() if not isinstance(val, str) else val.strip()
            else:
                rec.setdefault(field, None)
        if any(rec.get(f) for f in ('shot_num','scene_code','vfx_desc','script_desc','cost_est')):
            rows_out.append(rec)
    return rows_out, wb.sheetnames


def clean_cost(v, max_cost=_MAX_COST):
    """Sanitise a cost field. Returns None if value exceeds max_cost."""
    try:
        val = float(str(v).replace(',', '').replace('$', ''))
        if val < 0:
            return 0.0
        if val > max_cost:
            return None
        return val
    except Exception:
        return 0.0


def get_project_setting(key, default=None):
    """Read a per-project setting from the project_settings table."""
    try:
        conn = get_db()
        if not conn:
            return default
        row = conn.execute(
            'SELECT value FROM project_settings WHERE key=?', (key,)
        ).fetchone()
        if row is None:
            return default
        v = row['value']
        try:
            return float(v) if '.' in str(v) else int(v)
        except Exception:
            return v
    except Exception:
        return default
