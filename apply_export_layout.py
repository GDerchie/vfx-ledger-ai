"""
apply_export_layout.py
======================
Applies the EXPORT_EXCELL_LAYOUT.xlsx style to any generated Excel export file.
Drop this file next to START.bat and call it from your export pipeline.

Usage (standalone):
    python apply_export_layout.py <path_to_export.xlsx>

Usage (from Python):
    from apply_export_layout import apply_layout
    apply_layout("path/to/output.xlsx")

Style source: EXPORT_EXCELL_LAYOUT.xlsx  (ForDistribution sheet)
Extracted: 2026-03-18
"""

import sys
from openpyxl import load_workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.styles.numbers import FORMAT_CURRENCY_USD_SIMPLE
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# COLOUR PALETTE  (all derived from layout file)
# ─────────────────────────────────────────────
# Theme 9 (Office Accent6 = #70AD47 green)
BG_HEADER_DARK   = "C5DEB5"   # theme:9 tint:0.60 – row-2 header band / totals
BG_HEADER_LIGHT  = "E2EEDA"   # theme:9 tint:0.80 – column header rows (rows 3, 18)
BG_GREEN_SECTION = "E2EFD9"   # explicit hex – L/M/N columns (reductions/CTD/deadline)
BG_WHITE         = "FFFFFF"   # data rows

FONT_NAME        = "Calibri"
FONT_EMOJI       = "Segoe UI Emoji"
FONT_DARK_GREEN  = "548135"   # FF548135 – total $ amounts in green cols

# Number formats (exact from layout)
FMT_CURRENCY     = '_-[$\u0024-409]* #,##0.00_ ;_-[$\u0024-409]* \\-#,##0.00\\ ;_-[$\u0024-409]* "-"??_ ;_-@_ '
FMT_CURRENCY_ALT = '_-"$"* #,##0.00_-;\\-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@'

# ─────────────────────────────────────────────
# BORDER HELPERS
# ─────────────────────────────────────────────
def side(style="thin", color="FF000000"):
    return Side(style=style, color=color)

MEDIUM = side("medium")
DOTTED = side("dotted")
THIN   = side("thin")
NONE_  = Side(style=None)

def border(left=None, right=None, top=None, bottom=None):
    return Border(left=left or NONE_, right=right or NONE_,
                  top=top or NONE_,  bottom=bottom or NONE_)

# Common border combos used in the layout
B_OUTER    = border(MEDIUM, MEDIUM, MEDIUM, MEDIUM)
B_TOP_BOT  = border(top=MEDIUM, bottom=MEDIUM)
B_ROW_DATA = border(DOTTED, DOTTED, DOTTED, DOTTED)
B_LEFT_COL = border(left=MEDIUM, right=DOTTED, top=DOTTED, bottom=DOTTED)
B_RIGHT_CLOSE = border(left=DOTTED, right=MEDIUM, top=DOTTED, bottom=DOTTED)
B_SECTION_HEADER = border(MEDIUM, MEDIUM, MEDIUM, MEDIUM)

# ─────────────────────────────────────────────
# FILL HELPERS
# ─────────────────────────────────────────────
def fill(hex_rgb):
    return PatternFill(fill_type="solid", fgColor=hex_rgb)

FILL_HEADER_DARK  = fill(BG_HEADER_DARK)
FILL_HEADER_LIGHT = fill(BG_HEADER_LIGHT)
FILL_GREEN_SEC    = fill(BG_GREEN_SECTION)
FILL_WHITE        = fill(BG_WHITE)
FILL_NONE         = PatternFill(fill_type=None)

# ─────────────────────────────────────────────
# FONT HELPERS
# ─────────────────────────────────────────────
def font(bold=False, size=12, italic=False, color="FF000000", name=FONT_NAME):
    return Font(name=name, bold=bold, size=size, italic=italic, color=color)

# ─────────────────────────────────────────────
# ALIGNMENT HELPERS
# ─────────────────────────────────────────────
def align(h="general", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

# ─────────────────────────────────────────────
# COLUMN WIDTHS  (from layout file, cols A–O)
# ─────────────────────────────────────────────
COLUMN_WIDTHS = {
    "A":  4.00,
    "B":  7.00,
    "C":  6.25,
    "D": 18.08,
    "E":  5.58,
    "F": 10.75,
    "G": 12.08,
    "H": 18.33,
    "I": 16.83,
    "J": 15.91,
    "K": 21.83,
    "L": 42.66,
    "M": 34.25,
    "N": 22.00,
    "O": 11.16,
}

# ─────────────────────────────────────────────
# ROW HEIGHT DEFAULTS
# ─────────────────────────────────────────────
ROW_HEIGHT_DEFAULT = 15.75
ROW_HEIGHT_HEADER  = 26.50   # row 4 in layout (first data row of EFC summary)

# ─────────────────────────────────────────────
# CELL STYLE DEFINITIONS  (by semantic zone)
# ─────────────────────────────────────────────

def style_title_row(cell, text=None):
    """Row 2 – top banner: EFC SUMMARY, STATUS, TODAY()"""
    cell.font      = font(bold=True, size=12)
    cell.fill      = FILL_HEADER_DARK
    cell.alignment = align("left", "center")
    cell.border    = border(MEDIUM, DOTTED, MEDIUM, MEDIUM)
    if text is not None:
        cell.value = text

def style_section_label(cell):
    """B2, B15, B3 type labels – plain left-aligned titles"""
    cell.font      = font(bold=True, size=12)
    cell.alignment = align("left", "top")

def style_col_header(cell):
    """Row 3 / row 18 column headers – light green bg, bold, centered, wrap"""
    cell.font      = font(bold=True, size=10)
    cell.fill      = FILL_HEADER_LIGHT
    cell.alignment = align("center", "center", wrap=True)
    cell.border    = B_TOP_BOT

def style_col_header_green(cell):
    """L18/M18/N18 – green section header (reductions/CTD/deadline)"""
    cell.font      = font(bold=True, size=12)
    cell.fill      = FILL_GREEN_SEC
    cell.alignment = align("center", "center", wrap=True)
    cell.border    = B_OUTER

def style_ep_badge(cell, ep_number):
    """Column E episode number badge (EP 101, 102 …)"""
    cell.value     = ep_number
    cell.font      = font(bold=True, size=12)
    cell.alignment = align("center", "center")
    cell.border    = border(MEDIUM, DOTTED, DOTTED, DOTTED)

def style_data_cell(cell):
    """Standard data row cell (F–J, numeric data)"""
    cell.font      = font(size=12)
    cell.alignment = align("center", "center")
    cell.border    = B_ROW_DATA

def style_currency_cell(cell):
    """Budget / EFC / Variance columns (H, I, J)"""
    cell.font          = font(size=12)
    cell.alignment     = align("center", "center")
    cell.border        = B_ROW_DATA
    cell.number_format = FMT_CURRENCY

def style_status_cell(cell):
    """Column K – STATUS with emoji"""
    cell.font      = font(size=10, name=FONT_EMOJI)
    cell.alignment = align("center", "center", wrap=True)
    cell.border    = B_ROW_DATA

def style_notes_cell(cell):
    """Column L – Estimated Reductions / suggestions (italic, wrap)"""
    cell.font      = font(size=10, italic=True)
    cell.alignment = align("center", "center", wrap=True)
    cell.border    = B_ROW_DATA

def style_ctd_cell(cell):
    """Column M – Estimated CTD"""
    cell.font      = font(size=12)
    cell.alignment = align("center", "center", wrap=True)
    cell.border    = B_RIGHT_CLOSE

def style_deadline_cell(cell):
    """Column N – Turnover Deadline"""
    cell.font      = font(size=12)
    cell.alignment = align("center", "center", wrap=True)
    cell.border    = B_ROW_DATA

def style_total_row(cell, is_currency=False):
    """TOTAL summary row (row 12 / 225 / 435 pattern)"""
    cell.font  = font(bold=True, size=12)
    cell.fill  = FILL_HEADER_LIGHT
    cell.alignment = align("center", "center")
    cell.border = border(DOTTED, DOTTED, NONE_, MEDIUM)
    if is_currency:
        cell.number_format = FMT_CURRENCY

def style_script_version(cell, text="Script_V/Edit_V"):
    """Column D – Script/Edit version label (italic, right-aligned)"""
    cell.value     = text
    cell.font      = font(size=10, italic=True)
    cell.alignment = align("right", "center")
    cell.border    = border(right=MEDIUM, bottom=MEDIUM)

def style_today_cell(cell):
    """L2 TODAY() date cell"""
    cell.font          = font(bold=True, size=12, italic=True)
    cell.fill          = FILL_HEADER_DARK
    cell.alignment     = align("right", "center", wrap=True)
    cell.border        = border(DOTTED, DOTTED, MEDIUM, MEDIUM)
    cell.number_format = "DD-MMM-YYYY"

def style_version_cell(cell, version="v2.3"):
    """N1 version label"""
    cell.value     = version
    cell.font      = font(size=10)
    cell.alignment = align("right", "center")

# ─────────────────────────────────────────────
# MAIN APPLICATION FUNCTION
# ─────────────────────────────────────────────

def apply_layout(xlsx_path: str, sheet_name: str = None, version: str = "v2.3"):
    """
    Apply the ForDistribution layout style to an exported Excel file.

    Parameters
    ----------
    xlsx_path  : str   Path to the .xlsx file to style (modified in-place).
    sheet_name : str   Target worksheet name. If None, uses the active sheet.
    version    : str   Version string written to cell N1 (default 'v2.3').
    """
    wb = load_workbook(xlsx_path)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    print(f"[apply_layout] Styling sheet '{ws.title}' in '{xlsx_path}' ...")

    # ── 1. COLUMN WIDTHS ──────────────────────────────────────────────────
    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # ── 2. DEFAULT ROW HEIGHT ─────────────────────────────────────────────
    ws.sheet_format.defaultRowHeight = ROW_HEIGHT_DEFAULT
    ws.sheet_format.customHeight = True

    # ── 3. DETERMINE STRUCTURE  ───────────────────────────────────────────
    # Find max data row
    max_row = ws.max_row
    max_col = ws.max_column

    # ── 4. APPLY STYLES ROW BY ROW ────────────────────────────────────────

    for row in ws.iter_rows():
        row_num = row[0].row

        for cell in row:
            col_idx = cell.column    # 1-based
            col_ltr = get_column_letter(col_idx)

            # ---------- ROW 1: version label ----------
            if row_num == 1:
                if col_ltr == "N":
                    style_version_cell(cell, version)
                else:
                    cell.font = font(size=12)

            # ---------- ROW 2: EFC SUMMARY banner ----------
            elif row_num == 2:
                if col_ltr == "B":
                    style_section_label(cell)
                elif col_ltr in ("E", "F", "G", "H", "I", "J", "K", "L", "M"):
                    style_title_row(cell)
                    if col_ltr == "L":
                        style_today_cell(cell)
                elif col_ltr == "K":
                    cell.font      = font(bold=True, size=12)
                    cell.fill      = FILL_HEADER_DARK
                    cell.alignment = align("center", "center")
                    cell.border    = B_OUTER

            # ---------- ROW 3: EFC column headers ----------
            elif row_num == 3:
                if col_ltr in ("E", "F", "G", "H", "I", "J", "K", "L", "M"):
                    style_col_header(cell)
                    cell.fill = FILL_HEADER_LIGHT

            # ---------- ROW 4–11: per-episode summary rows ----------
            elif 4 <= row_num <= 11:
                ws.row_dimensions[row_num].height = 16.0
                if col_ltr == "D":
                    style_script_version(cell)
                elif col_ltr == "E":
                    style_ep_badge(cell, cell.value)
                elif col_ltr in ("F", "G"):
                    style_data_cell(cell)
                elif col_ltr in ("H", "I", "J"):
                    style_currency_cell(cell)
                elif col_ltr == "K":
                    style_status_cell(cell)
                elif col_ltr == "L":
                    style_notes_cell(cell)
                elif col_ltr == "M":
                    style_ctd_cell(cell)

            # ---------- ROW 12: EFC summary TOTAL ----------
            elif row_num == 12:
                if col_ltr in ("E", "K", "L", "M"):
                    style_total_row(cell, is_currency=False)
                    cell.fill = FILL_HEADER_LIGHT
                elif col_ltr in ("F", "G"):
                    style_total_row(cell, is_currency=False)
                    cell.fill = FILL_HEADER_LIGHT
                elif col_ltr in ("H", "I", "J"):
                    style_total_row(cell, is_currency=True)
                    cell.fill = FILL_HEADER_LIGHT

            # ---------- ROW 15: section title BREAKDOWN ----------
            elif row_num == 15:
                if col_ltr == "B":
                    cell.font = font(size=16)

            # ---------- ROW 17: group-span labels (SHOT COUNT / COST / ...) ----------
            elif row_num == 17:
                if col_ltr in ("F", "G"):
                    cell.font      = font(size=11)
                    cell.alignment = align("center", "center")
                    cell.border    = border(left=MEDIUM, top=MEDIUM)
                elif col_ltr in ("H", "I"):
                    cell.font      = font(size=11)
                    cell.alignment = align("center", "center")
                    cell.border    = border(left=MEDIUM, top=MEDIUM)
                elif col_ltr in ("L", "M", "N"):
                    cell.font      = font(size=11)
                    cell.alignment = align("center", "center")
                    cell.border    = border(left=MEDIUM, right=MEDIUM, top=MEDIUM)

            # ---------- EPISODE HEADER ROWS (18, 228, …) ----------
            elif cell.value in (
                "SEQ", "LOC", "OMIT", "ESTS#", "Current Cut#",
                "LBUDGET", "EFC", "VARIANCE", "STATUS",
                "ESTIMATED REDUCTIONS", "ESTIMATED CTD", "TURNOVER DEADLINE"
            ) or (col_ltr == "B" and isinstance(cell.value, int) and 100 <= cell.value <= 999):
                # Episode header band
                if col_ltr in ("B", "C", "D", "E", "F", "G", "H", "I", "J", "K"):
                    style_col_header(cell)
                elif col_ltr in ("L", "M", "N"):
                    style_col_header_green(cell)

            # ---------- DATA ROWS (seq breakdown rows) ----------
            else:
                ws.row_dimensions[row_num].height = ROW_HEIGHT_DEFAULT
                if col_ltr in ("B", "C", "D", "E"):
                    cell.font      = font(size=10)
                    cell.alignment = align("center", "center")
                    cell.border    = B_ROW_DATA
                elif col_ltr in ("F", "G"):
                    cell.font      = font(size=10)
                    cell.alignment = align("center", "center")
                    cell.border    = B_ROW_DATA
                elif col_ltr in ("H", "I"):
                    cell.font          = font(size=10)
                    cell.alignment     = align("center", "center")
                    cell.border        = B_ROW_DATA
                    cell.number_format = FMT_CURRENCY_ALT
                elif col_ltr == "J":
                    cell.font          = font(bold=True, size=10)
                    cell.alignment     = align("center", "center")
                    cell.border        = B_ROW_DATA
                    cell.number_format = FMT_CURRENCY
                elif col_ltr == "K":
                    style_status_cell(cell)
                elif col_ltr == "L":
                    cell.font      = font(size=10)
                    cell.alignment = align("center", "center", wrap=True)
                    cell.border    = B_ROW_DATA
                elif col_ltr == "M":
                    cell.font      = font(size=12, color="FF" + FONT_DARK_GREEN)
                    cell.alignment = align("center", "center", wrap=True)
                    cell.border    = B_RIGHT_CLOSE
                elif col_ltr == "N":
                    cell.font      = font(size=12)
                    cell.alignment = align("center", "center", wrap=True)
                    cell.border    = B_ROW_DATA

            # ---------- TOTAL rows (pattern: col C == "TOTAL") ----------
            # Catch totals rows regardless of row number
            if cell.value == "TOTAL" and col_ltr == "C":
                # Style the full total row
                for c in ws[row_num]:
                    cl = get_column_letter(c.column)
                    c.font  = font(bold=True, size=12)
                    c.fill  = FILL_HEADER_LIGHT
                    c.alignment = align("center", "center")
                    c.border = border(DOTTED, DOTTED, NONE_, MEDIUM)
                    if cl in ("H", "I"):
                        c.number_format = FMT_CURRENCY
                    if cl in ("J", "L", "M", "N"):
                        c.fill = FILL_GREEN_SEC

    # ── 5. PRINT AREA & FREEZE PANES ──────────────────────────────────────
    # Match layout: no freeze panes set
    # ws.freeze_panes = "C19"   # uncomment if you want to freeze header rows

    # ── 6. SHEET SETTINGS ─────────────────────────────────────────────────
    ws.sheet_view.showGridLines = True

    wb.save(xlsx_path)
    print(f"[apply_layout] Done. File saved: {xlsx_path}")
    return xlsx_path


# ─────────────────────────────────────────────
# CONVENIENCE WRAPPER for use in START.bat pipeline
# ─────────────────────────────────────────────

def apply_layout_to_latest_export(export_folder: str, pattern: str = "*.xlsx"):
    """
    Find the most-recently-modified .xlsx in export_folder and apply the layout.
    Useful when START.bat generates a timestamped filename each run.
    """
    import glob, os
    files = glob.glob(os.path.join(export_folder, pattern))
    if not files:
        raise FileNotFoundError(f"No .xlsx found in {export_folder}")
    latest = max(files, key=os.path.getmtime)
    return apply_layout(latest)


# ─────────────────────────────────────────────
# CLI ENTRY POINT
# ─────────────────────────────────────────────

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python apply_export_layout.py <path_to_export.xlsx> [sheet_name] [version]")
        sys.exit(1)

    path        = sys.argv[1]
    sheet       = sys.argv[2] if len(sys.argv) > 2 else None
    ver         = sys.argv[3] if len(sys.argv) > 3 else "v2.3"

    apply_layout(path, sheet_name=sheet, version=ver)
