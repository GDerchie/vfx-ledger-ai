"""routes/assets.py — Assets Blueprint."""
import io
import json
from flask import Blueprint, render_template, request, jsonify
import core

bp = Blueprint('assets', __name__)


@bp.route('/assets')
def assets_page():
    r = core.require_project()
    if r: return r
    return render_template('assets.html')


@bp.route('/api/assets', methods=['GET'])
def get_assets():
    conn = core.get_db()
    if not conn: return jsonify([])
    ep = request.args.get('ep', type=int)
    show_omit = request.args.get('omit', '0') == '1'
    q = 'SELECT * FROM assets'; params = []; conds = []
    if ep: conds.append('ep=?'); params.append(ep)
    if not show_omit: conds.append('omit=0')
    if conds: q += ' WHERE ' + ' AND '.join(conds)
    q += ' ORDER BY ep, orig_id, id'
    rows = conn.execute(q, params).fetchall()
    conn.close()
    result = []
    for r in rows:
        d = dict(r)
        try: d['vendor_bids'] = json.loads(d['vendor_bids'] or '{}')
        except: d['vendor_bids'] = {}
        result.append(d)
    return jsonify(result)


@bp.route('/api/assets/vendors', methods=['GET'])
def get_asset_vendors():
    conn = core.get_db()
    if not conn: return jsonify([])
    ep = request.args.get('ep', type=int)
    q = 'SELECT vendor_bids FROM assets WHERE omit=0'; params = []
    if ep: q += ' AND ep=?'; params.append(ep)
    rows = conn.execute(q, params).fetchall()
    conn.close()
    vendors, seen = [], set()
    for r in rows:
        try:
            for v in json.loads(r['vendor_bids'] or '{}').keys():
                if v not in seen: seen.add(v); vendors.append(v)
        except: pass
    return jsonify(vendors)


@bp.route('/api/assets/orphaned')
def orphaned_assets():
    conn = core.get_db()
    if not conn: return jsonify([])
    rows = conn.execute('''
        SELECT a.* FROM assets a
        LEFT JOIN shots s ON UPPER(TRIM(s.asset)) LIKE '%'||UPPER(TRIM(a.asset_name))||'%'
        WHERE s.id IS NULL AND a.omit=0
    ''').fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@bp.route('/api/assets/summary', methods=['GET'])
def get_assets_summary():
    conn = core.get_db()
    if not conn: return jsonify([])
    rows = conn.execute('''
        SELECT ep,
               COUNT(CASE WHEN omit=0 THEN 1 END) AS asset_count,
               SUM(CASE WHEN omit=0 THEN COALESCE(est_budget,0) ELSE 0 END) AS est_budget,
               SUM(CASE WHEN omit=0 THEN COALESCE(actual_spend,0) ELSE 0 END) AS actual_spend
        FROM assets GROUP BY ep ORDER BY ep
    ''').fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@bp.route('/api/assets/<int:asset_id>', methods=['PUT'])
def update_asset(asset_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    d = request.json
    allowed = ['award_vendor', 'est_budget', 'actual_spend', 'notes', 'omit', 'turnover']
    updates = {k: d[k] for k in allowed if k in d}
    for f in ('est_budget', 'actual_spend'):
        if f in updates:
            cleaned = core.clean_cost(updates[f])
            if cleaned is None:
                return jsonify({'error': f'Value for {f} exceeds maximum allowed (${core._MAX_COST:,.0f})'}), 400
            updates[f] = cleaned
    if updates:
        sql = 'UPDATE assets SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [asset_id])
        conn.commit()
    conn.close()
    core.cache.clear()
    return jsonify({'status': 'ok'})


@bp.route('/api/assets/import', methods=['POST'])
def import_assets():
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    f = request.files.get('file')
    if not f: return jsonify({'error': 'No file'}), 400
    if not f.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Use .xlsx files'}), 400
    mode = request.form.get('mode', 'replace')
    ep_filter = request.form.get('ep', type=int)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    ws = wb.active if 'Assets' not in wb.sheetnames else wb['Assets']
    header_row_idx = None; headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [str(v).strip() if v is not None else '' for v in row]
        if 'ASSET NAME' in vals or 'ITEM' in vals:
            header_row_idx = i; headers = vals; break
    if header_row_idx is None:
        return jsonify({'error': 'Could not find header row (needs ASSET NAME column)'}), 400
    col_map = {}; vendor_cols = []
    fixed = {'ITEM':'orig_id','EP':'ep','SCENE #':'scene_code','TURNOVER':'turnover',
             'ASSET NAME':'asset_name','DESCRIPTION':'description','REFERENCE / IMAGE':'reference',
             'ASSET TYPE (ENVIRO/CHAR/CREATURE/ PROP/WEAPON, ETC)':'asset_type','ASSET TYPE':'asset_type',
             'REPEAT ASSET(SEASON ASSET)':'repeat_asset','REPEAT ASSET':'repeat_asset',
             'HERO ID':'hero_id','OMIT':'omit','REF':'ref','NOTES':'notes',
             'LIDAR':'lidar','PHOTOGRAMATRY':'photogrammetry','PHOTOGRAMMETRY':'photogrammetry',
             'ESTBUDGET':'est_budget','ACTUAL SPEND':'actual_spend','AWARD VENDOR':'award_vendor'}
    for i, h in enumerate(headers):
        hu = h.strip().upper()
        if hu in fixed: col_map[fixed[hu]] = i
        elif hu and not hu.startswith('VARIANCE') and i > col_map.get('award_vendor', -1):
            if 'award_vendor' in col_map and i > col_map['award_vendor'] + 1:
                vendor_cols.append((i, h.strip()))
    if mode == 'replace':
        if ep_filter: conn.execute('DELETE FROM assets WHERE ep=?', (ep_filter,))
        else: conn.execute('DELETE FROM assets')
    def gv(field, row, default=None):
        idx = col_map.get(field)
        return row[idx] if idx is not None and idx < len(row) else default
    def flt(v, d=0.0):
        try: return float(v) if v is not None else d
        except: return d
    def nt(v, d=0):
        try: return int(float(v)) if v is not None else d
        except: return d
    def bl(v): return 1 if v and str(v).strip().lower() in ('true','1','yes','x') else 0
    inserted = 0
    for r in range(header_row_idx + 1, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if all(v is None for v in row): continue
        ep_val = gv('ep', row)
        if ep_val is None: continue
        try: ep_int = int(ep_val)
        except: continue
        if ep_filter and ep_int != ep_filter: continue
        bids = {}
        for vi, vname in vendor_cols:
            val = row[vi] if vi < len(row) else None
            if val is not None and str(val).strip() not in ('', '---'):
                try: bids[vname] = float(val)
                except: pass
        sc = gv('scene_code', row); asset_name = gv('asset_name', row)
        if sc is None and asset_name is None: continue
        conn.execute('''INSERT INTO assets
            (orig_id, ep, scene_code, turnover, asset_name, description, reference,
             asset_type, repeat_asset, hero_id, omit, ref, notes, lidar, photogrammetry,
             est_budget, actual_spend, award_vendor, vendor_bids) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (nt(gv('orig_id', row)), ep_int,
             str(sc).strip() if sc is not None else '',
             str(gv('turnover', row, '') or '').strip(),
             str(asset_name or '').strip(),
             str(gv('description', row, '') or '').strip(),
             str(gv('reference', row, '') or '').strip(),
             str(gv('asset_type', row, '') or '').strip(),
             bl(gv('repeat_asset', row)), bl(gv('hero_id', row)), bl(gv('omit', row)),
             str(gv('ref', row, '') or '').strip(), str(gv('notes', row, '') or '').strip(),
             bl(gv('lidar', row)), bl(gv('photogrammetry', row)),
             flt(gv('est_budget', row)), flt(gv('actual_spend', row)),
             str(gv('award_vendor', row, '') or '').strip(), json.dumps(bids)))
        inserted += 1
    conn.commit(); conn.close()
    return jsonify({'inserted': inserted})


@bp.route('/api/assets/clear', methods=['DELETE'])
def clear_assets():
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    ep = request.args.get('ep', type=int)
    if ep: conn.execute('DELETE FROM assets WHERE ep=?', (ep,))
    else: conn.execute('DELETE FROM assets')
    conn.commit(); conn.close()
    return jsonify({'status': 'ok'})


# ── Asset reuse tracking ──────────────────────────────────────────────────────

@bp.route('/api/assets/<int:asset_id>/shots', methods=['GET'])
def asset_shots(asset_id):
    """List shots linked to a given asset."""
    conn = core.get_db()
    if not conn: return jsonify([])
    rows = conn.execute('''
        SELECT s.id, s.ep, s.shot_num, s.scene_code, s.location,
               s.script_desc, s.vfx_desc, s.award_vendor
        FROM shot_asset_links sal
        JOIN shots s ON s.id = sal.shot_id
        WHERE sal.asset_id = ?
        ORDER BY s.ep, s.shot_num
    ''', (asset_id,)).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@bp.route('/api/shot_asset_links', methods=['POST'])
def create_shot_asset_link():
    """Link a shot to an asset (idempotent)."""
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    d = request.json or {}
    shot_id = d.get('shot_id')
    asset_id = d.get('asset_id')
    if not shot_id or not asset_id:
        return jsonify({'error': 'shot_id and asset_id required'}), 400
    conn.execute(
        'INSERT OR IGNORE INTO shot_asset_links (shot_id, asset_id) VALUES (?,?)',
        (shot_id, asset_id)
    )
    conn.commit(); conn.close()
    return jsonify({'status': 'ok'})


@bp.route('/api/shot_asset_links', methods=['DELETE'])
def delete_shot_asset_link():
    """Unlink a shot from an asset."""
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    d = request.json or {}
    shot_id = d.get('shot_id')
    asset_id = d.get('asset_id')
    if not shot_id or not asset_id:
        return jsonify({'error': 'shot_id and asset_id required'}), 400
    conn.execute(
        'DELETE FROM shot_asset_links WHERE shot_id=? AND asset_id=?',
        (shot_id, asset_id)
    )
    conn.commit(); conn.close()
    return jsonify({'status': 'ok'})


@bp.route('/api/assets/reuse_summary', methods=['GET'])
def asset_reuse_summary():
    """Return assets with their shot-link counts, sorted by reuse descending."""
    conn = core.get_db()
    if not conn: return jsonify([])
    ep = request.args.get('ep', type=int)
    q = '''
        SELECT a.id, a.ep, a.asset_name, a.asset_type, a.award_vendor,
               COUNT(sal.shot_id) AS shot_count
        FROM assets a
        LEFT JOIN shot_asset_links sal ON sal.asset_id = a.id
        WHERE a.omit = 0
    '''
    params = []
    if ep:
        q += ' AND a.ep = ?'; params.append(ep)
    q += ' GROUP BY a.id ORDER BY shot_count DESC, a.ep, a.asset_name'
    rows = conn.execute(q, params).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])
