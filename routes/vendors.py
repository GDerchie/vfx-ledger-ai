"""routes/vendors.py — Vendor tracker, bid compare, invoice log, scenario, forecast Blueprint."""
import io
import json
from statistics import mean, stdev
from flask import Blueprint, render_template, request, jsonify
import core

bp = Blueprint('vendors', __name__)


# ── Page routes ──────────────────────────────────────────────────────────────

@bp.route('/bidcompare')
def bidcompare_page():
    r = core.require_project()
    if r: return r
    return render_template('bidcompare.html')


@bp.route('/vendortracker')
def vendortracker_page():
    r = core.require_project()
    if r: return r
    return render_template('vendortracker.html')


@bp.route('/invoicelog')
def invoicelog_page():
    r = core.require_project()
    if r: return r
    return render_template('invoicelog.html')


@bp.route('/scenario')
def scenario():
    r = core.require_project()
    if r: return r
    return render_template('scenario.html')


# ── API – Bid Compare ────────────────────────────────────────────────────────

@bp.route('/api/bidcompare', methods=['GET'])
def get_bidcompare():
    conn = core.get_db()
    if not conn: return jsonify([])
    ep = request.args.get('ep', type=int)
    show_novfx = request.args.get('novfx', '0') == '1'
    if ep:
        if show_novfx:
            rows = conn.execute('SELECT * FROM bid_compare WHERE ep=? ORDER BY id', (ep,)).fetchall()
        else:
            rows = conn.execute('SELECT * FROM bid_compare WHERE ep=? AND novfx=0 ORDER BY id', (ep,)).fetchall()
    else:
        if show_novfx:
            rows = conn.execute('SELECT * FROM bid_compare ORDER BY ep, id').fetchall()
        else:
            rows = conn.execute('SELECT * FROM bid_compare WHERE novfx=0 ORDER BY ep, id').fetchall()
    conn.close()
    result = []
    for r in rows:
        d = dict(r)
        try: d['vendor_bids'] = json.loads(d['vendor_bids'] or '{}')
        except: d['vendor_bids'] = {}
        result.append(d)
    return jsonify(result)


@bp.route('/api/bidcompare/vendors', methods=['GET'])
def get_bidcompare_vendors():
    conn = core.get_db()
    if not conn: return jsonify([])
    ep = request.args.get('ep', type=int)
    if ep:
        rows = conn.execute('SELECT vendor_bids FROM bid_compare WHERE ep=? AND novfx=0', (ep,)).fetchall()
    else:
        rows = conn.execute('SELECT vendor_bids FROM bid_compare WHERE novfx=0').fetchall()
    conn.close()
    vendors = []; seen = set()
    for r in rows:
        try:
            bids = json.loads(r['vendor_bids'] or '{}')
            for v in bids.keys():
                if v not in seen: seen.add(v); vendors.append(v)
        except: pass
    return jsonify(vendors)


@bp.route('/api/bidcompare/import', methods=['POST'])
def import_bidcompare():
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    f = request.files.get('file')
    if not f: return jsonify({'error': 'No file'}), 400
    if not f.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Use .xlsx files'}), 400
    mode = request.form.get('mode', 'replace')
    ep_filter = request.form.get('ep', type=int)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    ws = wb.active
    header_row_idx = None; headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [str(v).strip() if v is not None else '' for v in row]
        if 'ID' in vals and 'EP' in vals:
            header_row_idx = i; headers = vals; break
    if header_row_idx is None:
        return jsonify({'error': 'Could not find header row (needs ID and EP columns)'}), 400
    col_map = {}; vendor_cols = []
    fixed = {'ID':'orig_id','EP':'ep','SC':'sc','SETTING':'setting','NOVFX':'novfx',
             'SCOUNT':'scount','EDIT COUNT':'edit_count','VFXTYPE':'vfxtype',
             'VERSION':'version','AWARD':'award','LOCKBUDGET':'lockbudget','EFC':'efc'}
    for i, h in enumerate(headers):
        hu = h.strip().upper()
        if hu in fixed: col_map[fixed[hu]] = i
        elif hu and hu not in ('', 'NONE') and i > max(col_map.values(), default=-1):
            if 'efc' in col_map and i > col_map['efc']:
                vendor_cols.append((i, h.strip()))
    if mode == 'replace':
        if ep_filter: conn.execute('DELETE FROM bid_compare WHERE ep=?', (ep_filter,))
        else: conn.execute('DELETE FROM bid_compare')
    inserted = 0
    for r in range(header_row_idx + 1, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if all(v is None for v in row): continue
        ep_val = row[col_map['ep']] if 'ep' in col_map and col_map['ep'] < len(row) else None
        if ep_val is None: continue
        try: ep_int = int(ep_val)
        except: continue
        if ep_filter and ep_int != ep_filter: continue

        def gv(field, default=None):
            idx = col_map.get(field)
            if idx is None or idx >= len(row): return default
            return row[idx]

        def flt(v, default=0.0):
            try: return float(v) if v is not None and str(v) != '#REF!' else default
            except: return default

        def nt(v):
            try: return int(float(v)) if v is not None else 0
            except: return 0

        bids = {}
        for vi, vname in vendor_cols:
            val = row[vi] if vi < len(row) else None
            if val is not None and str(val).strip() not in ('', '---', 'None'):
                try: bids[vname] = float(val)
                except: pass
        novfx_val = gv('novfx')
        novfx = 1 if novfx_val and str(novfx_val).strip().lower() in ('true','1','yes') else 0
        conn.execute('''INSERT INTO bid_compare
            (orig_id, ep, sc, setting, novfx, scount, edit_count, vfxtype,
             version, award, lockbudget, efc, vendor_bids)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (nt(gv('orig_id')), ep_int,
             str(gv('sc', '')).strip() if gv('sc') is not None else '',
             str(gv('setting', '')).strip() if gv('setting') is not None else '',
             novfx, nt(gv('scount')), nt(gv('edit_count')),
             str(gv('vfxtype', '')).strip() if gv('vfxtype') is not None else '',
             nt(gv('version', 1)),
             str(gv('award', '')).strip() if gv('award') is not None else '',
             flt(gv('lockbudget')), flt(gv('efc')),
             json.dumps(bids)))
        inserted += 1
    conn.commit(); conn.close()
    return jsonify({'inserted': inserted})


@bp.route('/api/bidcompare/<int:bid_id>', methods=['PUT'])
def update_bidcompare(bid_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    d = request.json
    allowed = ['award', 'lockbudget', 'efc', 'notes']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE bid_compare SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [bid_id])
        conn.commit()
    conn.close()
    return jsonify({'status': 'ok'})


@bp.route('/api/bidcompare/clear', methods=['DELETE'])
def clear_bidcompare():
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    ep = request.args.get('ep', type=int)
    if ep: conn.execute('DELETE FROM bid_compare WHERE ep=?', (ep,))
    else: conn.execute('DELETE FROM bid_compare')
    conn.commit(); conn.close()
    return jsonify({'status': 'ok'})


# ── API – Vendor Tracker ─────────────────────────────────────────────────────

@bp.route('/api/vendortracker', methods=['GET', 'POST'])
def vendortracker_collection():
    conn = core.get_db()
    if not conn:
        return (jsonify([]) if request.method == 'GET' else jsonify({'error': 'No project'})), (200 if request.method == 'GET' else 400)
    if request.method == 'GET':
        rows = conn.execute('SELECT * FROM vendor_tracker ORDER BY ep, vendor').fetchall()
        conn.close()
        return jsonify([dict(r) for r in rows])
    d = request.json; c = conn.cursor()
    c.execute('''INSERT INTO vendor_tracker
        (vendor, region, ep, asset_award, award_ep, tot_award,
         paid, pending, remaining,
         tax_reb, gross_local, gross_usd, sale_tax_amount,
         gross_plus_tax, tax_rebate_amt, net_after_rebate, final_price)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
        (d.get('vendor',''), d.get('region',''),
         d.get('ep'), d.get('asset_award'), d.get('award_ep'), d.get('tot_award'),
         d.get('paid', 0), d.get('pending', 0), d.get('remaining'),
         d.get('tax_reb'), d.get('gross_local'), d.get('gross_usd'),
         d.get('sale_tax_amount'), d.get('gross_plus_tax'),
         d.get('tax_rebate_amt'), d.get('net_after_rebate'), d.get('final_price')))
    new_id = c.lastrowid
    conn.commit(); conn.close()
    return jsonify({'id': new_id})


@bp.route('/api/vendortracker/<int:row_id>', methods=['PUT'])
def update_vendortracker(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    d = request.json
    allowed = ['vendor', 'region', 'ep', 'asset_award', 'award_ep', 'tot_award',
               'paid', 'pending', 'remaining', 'tax_reb', 'gross_local', 'gross_usd',
               'sale_tax_amount', 'gross_plus_tax', 'tax_rebate_amt', 'net_after_rebate', 'final_price']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE vendor_tracker SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [row_id])
        conn.commit()
    conn.close()
    return jsonify({'status': 'ok'})


@bp.route('/api/vendortracker/<int:row_id>', methods=['DELETE'])
def delete_vendortracker(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    conn.execute('DELETE FROM vendor_tracker WHERE id=?', (row_id,))
    conn.commit(); conn.close()
    return jsonify({'status': 'ok'})


@bp.route('/api/vendortracker/clear', methods=['DELETE'])
def clear_vendortracker():
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    conn.execute('DELETE FROM vendor_tracker')
    conn.execute('DELETE FROM invoice_log')
    conn.commit(); conn.close()
    return jsonify({'status': 'ok'})


@bp.route('/api/vendortracker/sync-from-invoices', methods=['POST'])
def sync_vendortracker_from_invoices():
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    pairs = conn.execute(
        'SELECT DISTINCT vendor, episode FROM invoice_log WHERE vendor IS NOT NULL AND episode IS NOT NULL'
    ).fetchall()
    for row in pairs:
        core.sync_tracker_paid(conn, row['vendor'], row['episode'])
    conn.commit(); conn.close()
    return jsonify({'synced': len(pairs)})


@bp.route('/api/vendortracker/import', methods=['POST'])
def import_vendortracker():
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    f = request.files.get('file')
    if not f: return jsonify({'error': 'No file'}), 400
    if not f.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Use .xlsx files'}), 400
    mode = request.form.get('mode', 'replace')
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    tracker_inserted = 0; invoice_inserted = 0

    # ── VENDORINVTRACK sheet ─────────────────────────────────────────────────
    if 'VENDORINVTRACK' in wb.sheetnames:
        ws = wb['VENDORINVTRACK']
        rows = list(ws.iter_rows(values_only=True))
        track_row = None
        for i, row in enumerate(rows):
            vals = [str(v).strip().upper() if v else '' for v in row]
            if 'VENDOR' in vals and 'PAID' in vals and 'PENDING' in vals:
                track_row = i; break
        if track_row is not None:
            headers = [str(v).strip() if v else '' for v in rows[track_row]]
            hmap = {}
            for i, h in enumerate(headers):
                hu = h.upper()
                if hu and hu not in hmap: hmap[hu] = i
            vendor_indices = [i for i, h in enumerate(headers) if h.strip().upper() == 'VENDOR']
            right_offset = vendor_indices[1] if len(vendor_indices) > 1 else None
            if mode == 'replace':
                conn.execute('DELETE FROM vendor_tracker')
            for row in rows[track_row + 1:]:
                if all(v is None for v in row): continue
                v_idx = hmap.get('VENDOR', 1)
                vendor = row[v_idx] if v_idx < len(row) else None
                if not vendor or str(vendor).strip() == '': continue
                vendor = str(vendor).strip()
                reg_idx = hmap.get('REG', 2)
                ep_idx  = hmap.get('EP', 3)
                ep_val  = row[ep_idx] if ep_idx < len(row) else None
                try:    ep_int = int(ep_val) if ep_val is not None else None
                except: ep_int = None

                def gv(key, default=None):
                    idx = hmap.get(key.upper())
                    return row[idx] if idx is not None and idx < len(row) else default

                def gv_right(offset, default=None):
                    if right_offset is None: return default
                    idx = right_offset + offset
                    return row[idx] if idx < len(row) else default

                tot_key = next((k for k in hmap if 'TOT' in k and 'AWARD' in k), None)
                rem_key = next((k for k in hmap if 'REMAIN' in k), None)
                conn.execute('''INSERT INTO vendor_tracker
                    (vendor, region, ep, asset_award, award_ep, tot_award,
                     paid, pending, remaining,
                     tax_reb, gross_local, gross_usd, sale_tax_amount,
                     gross_plus_tax, tax_rebate_amt, net_after_rebate, final_price)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                    (vendor,
                     str(row[reg_idx]).strip() if reg_idx < len(row) and row[reg_idx] else '',
                     ep_int,
                     core._sfloat(gv('ASSET_AWARD')),
                     core._sfloat(gv('AWARD_EP')),
                     core._sfloat(gv(tot_key) if tot_key else None),
                     core._sfloat(gv('PAID')) or 0,
                     core._sfloat(gv('PENDING')) or 0,
                     core._sfloat(gv(rem_key) if rem_key else None),
                     core._sfloat(gv_right(2)),
                     core._sfloat(gv_right(3)),
                     core._sfloat(gv_right(4)),
                     core._sfloat(gv_right(5)),
                     core._sfloat(gv_right(6)),
                     core._sfloat(gv_right(7)),
                     core._sfloat(gv_right(8)),
                     core._sfloat(gv_right(10)),
                    ))
                tracker_inserted += 1

    # ── INVOICELOG sheet ─────────────────────────────────────────────────────
    if 'INVOICELOG' in wb.sheetnames:
        ws2 = wb['INVOICELOG']
        rows2 = list(ws2.iter_rows(values_only=True))
        hdr_row = None
        for i, row in enumerate(rows2):
            vals = [str(v).strip().upper() if v else '' for v in row]
            if 'VENDOR' in vals and 'AMOUNT' in vals:
                hdr_row = i; break
        if hdr_row is not None:
            headers2 = [str(v).strip().upper() if v else '' for v in rows2[hdr_row]]
            h2map = {h: i for i, h in enumerate(headers2) if h}
            if mode == 'replace':
                conn.execute('DELETE FROM invoice_log')

            def g2(key, row):
                idx = h2map.get(key)
                return row[idx] if idx is not None and idx < len(row) else None

            def fmt_date(v):
                if v is None: return None
                if hasattr(v, 'strftime'): return v.strftime('%Y-%m-%d')
                return str(v).strip()

            for row in rows2[hdr_row + 1:]:
                if all(v is None for v in row): continue
                vendor = g2('VENDOR', row)
                if not vendor or str(vendor).strip() == '': continue
                conn.execute('''INSERT INTO invoice_log
                    (vendor, episode, inv_num, inv_date, amount, status, approve_date, notes)
                    VALUES (?,?,?,?,?,?,?,?)''',
                    (str(vendor).strip(),
                     int(g2('EPISODE', row)) if g2('EPISODE', row) is not None else None,
                     str(g2('INV#', row) or '').strip(),
                     fmt_date(g2('INV_DATE', row)),
                     core._sfloat(g2('AMOUNT', row)) or 0,
                     str(g2('STATUS', row) or '').strip(),
                     fmt_date(g2('APPROVE DATE', row)),
                     str(g2('NOTES', row) or '').strip()))
                invoice_inserted += 1

    conn.commit(); conn.close()
    return jsonify({'tracker_inserted': tracker_inserted, 'invoice_inserted': invoice_inserted})


# ── API – Invoice Log ────────────────────────────────────────────────────────

@bp.route('/api/invoicelog', methods=['GET', 'POST'])
def invoicelog_collection():
    conn = core.get_db()
    if not conn:
        return (jsonify([]) if request.method == 'GET' else jsonify({'error': 'No project'})), (200 if request.method == 'GET' else 400)
    if request.method == 'GET':
        rows = conn.execute('SELECT * FROM invoice_log ORDER BY episode, id').fetchall()
        conn.close()
        return jsonify([dict(r) for r in rows])
    d = request.json; c = conn.cursor()
    c.execute('''INSERT INTO invoice_log
        (vendor, episode, inv_num, inv_date, amount, status, approve_date, notes)
        VALUES (?,?,?,?,?,?,?,?)''',
        (d.get('vendor',''), d.get('episode'), d.get('inv_num',''),
         d.get('inv_date',''), d.get('amount', 0), d.get('status',''),
         d.get('approve_date',''), d.get('notes','')))
    new_id = c.lastrowid
    core.sync_tracker_paid(conn, d.get('vendor',''), d.get('episode'))
    conn.commit(); conn.close()
    return jsonify({'id': new_id})


@bp.route('/api/invoicelog/<int:row_id>', methods=['PUT'])
def update_invoicelog(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    old = conn.execute('SELECT vendor, episode FROM invoice_log WHERE id=?', (row_id,)).fetchone()
    d = request.json
    allowed = ['vendor', 'episode', 'inv_num', 'inv_date', 'amount', 'status', 'approve_date', 'notes']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE invoice_log SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [row_id])
    new_vendor  = updates.get('vendor',  old['vendor']  if old else None)
    new_episode = updates.get('episode', old['episode'] if old else None)
    if old:
        core.sync_tracker_paid(conn, old['vendor'], old['episode'])
    if new_vendor != (old['vendor'] if old else None) or new_episode != (old['episode'] if old else None):
        core.sync_tracker_paid(conn, new_vendor, new_episode)
    conn.commit(); conn.close()
    if old and 'status' in updates:
        core.log_change('invoice_log', row_id, 'status', old.get('status', ''), updates['status'])
        if str(updates['status']).lower() == 'approved':
            core.slack_alert(f'[VFX Budget] Invoice APPROVED \u2014 #{updates.get("inv_num","?")} {new_vendor} EP{new_episode} ${updates.get("amount",0):,.0f}')
    if old and 'amount' in updates:
        core.log_change('invoice_log', row_id, 'amount', old.get('amount', 0), updates['amount'])
    return jsonify({'status': 'ok'})


@bp.route('/api/invoicelog/<int:row_id>', methods=['DELETE'])
def delete_invoicelog(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    old = conn.execute('SELECT vendor, episode FROM invoice_log WHERE id=?', (row_id,)).fetchone()
    conn.execute('DELETE FROM invoice_log WHERE id=?', (row_id,))
    if old:
        core.sync_tracker_paid(conn, old['vendor'], old['episode'])
    conn.commit(); conn.close()
    return jsonify({'status': 'ok'})


# ── API – Vendor Registry ────────────────────────────────────────────────────

@bp.route('/api/vendor_registry', methods=['GET', 'POST'])
def vendor_registry_collection():
    conn = core.get_db()
    if not conn:
        return (jsonify([]) if request.method == 'GET' else jsonify({'error': 'No project'})), (200 if request.method == 'GET' else 400)
    if request.method == 'GET':
        rows = conn.execute('SELECT * FROM vendor_registry ORDER BY vendor').fetchall()
        conn.close()
        return jsonify([dict(r) for r in rows])
    d = request.json
    vendor = (d.get('vendor') or '').strip()
    if not vendor:
        conn.close(); return jsonify({'error': 'vendor required'}), 400
    c = conn.cursor()
    c.execute('''INSERT OR IGNORE INTO vendor_registry (vendor, region, fx_rate, tax_pct, rebate_pct, contact)
                 VALUES (?,?,?,?,?,?)''',
              (vendor, d.get('region',''), d.get('fx_rate', 1.0),
               d.get('tax_pct', 0.0), d.get('rebate_pct', 0.0), d.get('contact','')))
    new_id = c.lastrowid
    conn.commit(); conn.close()
    return jsonify({'id': new_id})


@bp.route('/api/vendor_registry/<int:row_id>', methods=['PUT', 'DELETE'])
def vendor_registry_row(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    if request.method == 'DELETE':
        conn.execute('DELETE FROM vendor_registry WHERE id=?', (row_id,))
        conn.commit(); conn.close(); return jsonify({'status': 'ok'})
    d = request.json
    allowed = ['vendor', 'region', 'fx_rate', 'tax_pct', 'rebate_pct', 'contact']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE vendor_registry SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [row_id])
        conn.commit()
    conn.close(); return jsonify({'status': 'ok'})


# ── API – Vendor Capacity ────────────────────────────────────────────────────

@bp.route('/api/vendor_capacity', methods=['GET', 'POST'])
def vendor_capacity_collection():
    conn = core.get_db()
    if not conn:
        return (jsonify([]) if request.method == 'GET' else jsonify({'error': 'No project'})), (200 if request.method == 'GET' else 400)
    if request.method == 'GET':
        rows = conn.execute('SELECT * FROM vendor_capacity ORDER BY vendor').fetchall()
        conn.close(); return jsonify([dict(r) for r in rows])
    d = request.json
    vendor = (d.get('vendor') or '').strip()
    if not vendor:
        conn.close(); return jsonify({'error': 'vendor required'}), 400
    c = conn.cursor()
    c.execute('''INSERT OR IGNORE INTO vendor_capacity (vendor, shots_per_month, efficiency_pct)
                 VALUES (?,?,?)''',
              (vendor, d.get('shots_per_month', 100), d.get('efficiency_pct', 0.95)))
    new_id = c.lastrowid
    conn.commit(); conn.close(); return jsonify({'id': new_id})


@bp.route('/api/vendor_capacity/<int:row_id>', methods=['PUT', 'DELETE'])
def vendor_capacity_row(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    if request.method == 'DELETE':
        conn.execute('DELETE FROM vendor_capacity WHERE id=?', (row_id,))
        conn.commit(); conn.close(); return jsonify({'status': 'ok'})
    d = request.json
    allowed = ['vendor', 'shots_per_month', 'efficiency_pct']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE vendor_capacity SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [row_id])
        conn.commit()
    conn.close(); return jsonify({'status': 'ok'})


# ── API – Vendor Loads (capacity utilization) ────────────────────────────────

@bp.route('/api/vendor_loads')
def vendor_loads():
    """Return per-vendor shot load vs capacity for the season."""
    conn = core.get_db()
    if not conn: return jsonify([])
    loads = core.get_vendor_loads(conn)
    conn.close()
    return jsonify(sorted(loads.values(), key=lambda x: -(x['assigned'] or 0)))


# ── API – Vendor Confidence ──────────────────────────────────────────────────

@bp.route('/api/vendor_confidence')
def vendor_confidence():
    conn = core.get_db()
    if not conn: return jsonify({})
    rows = conn.execute('SELECT vendor_bids, vfxtype FROM bid_compare WHERE novfx=0').fetchall()
    conn.close()
    buckets = {}
    for r in rows:
        try: bids = json.loads(r['vendor_bids'] or '{}')
        except: bids = {}
        for vendor, amount in bids.items():
            if amount is None: continue
            try: amount = float(amount)
            except: continue
            buckets.setdefault(vendor, []).append(amount)
    result = {}
    for vendor, amounts in buckets.items():
        n = len(amounts)
        if n < 2:
            cv = 1.0
        else:
            m = mean(amounts)
            sd = stdev(amounts)
            cv = (sd / m) if m else 1.0
        sample_factor = min(n / 30.0, 1.0)
        stability     = max(0.0, 1.0 - cv)
        confidence    = round(sample_factor * stability, 3)
        result[vendor] = {
            'samples':    n,
            'mean':       round(mean(amounts), 0) if amounts else 0,
            'cv':         round(cv, 3),
            'confidence': confidence,
            'label':      'HIGH' if confidence >= 0.7 else ('MED' if confidence >= 0.4 else 'LOW'),
        }
    return jsonify(result)


# ── API – Invoice-Shot Links ─────────────────────────────────────────────────

@bp.route('/api/invoicelog/<int:invoice_id>/shots', methods=['GET'])
def invoice_shot_list(invoice_id):
    """Return shots linked to this invoice, plus all shots for the same EP (for the modal)."""
    conn = core.get_db()
    if not conn: return jsonify({'linked': [], 'shots': []})
    linked_ids = {r['shot_id'] for r in conn.execute(
        'SELECT shot_id FROM invoice_shot_links WHERE invoice_id=?', (invoice_id,)
    ).fetchall()}
    inv = conn.execute('SELECT episode FROM invoice_log WHERE id=?', (invoice_id,)).fetchone()
    ep  = inv['episode'] if inv else None
    shots = []
    if ep:
        rows = conn.execute(
            'SELECT id, scene_code, s_code, shot_type, complexity, award_vendor, cost_est '
            'FROM shots WHERE ep=? AND omit=0 ORDER BY scene_code, s_code', (ep,)
        ).fetchall()
        shots = [dict(r) for r in rows]
    conn.close()
    return jsonify({'linked': list(linked_ids), 'shots': shots})


@bp.route('/api/invoicelog/<int:invoice_id>/shots', methods=['POST'])
def invoice_shot_link(invoice_id):
    """Link a shot to an invoice. Body: {shot_id: int}"""
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    shot_id = (request.json or {}).get('shot_id')
    if not shot_id:
        conn.close(); return jsonify({'error': 'shot_id required'}), 400
    conn.execute(
        'INSERT OR IGNORE INTO invoice_shot_links (invoice_id, shot_id) VALUES (?,?)',
        (invoice_id, shot_id)
    )
    conn.commit(); conn.close()
    return jsonify({'status': 'linked'})


@bp.route('/api/invoicelog/<int:invoice_id>/shots/<int:shot_id>', methods=['DELETE'])
def invoice_shot_unlink(invoice_id, shot_id):
    """Remove a shot→invoice link."""
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    conn.execute(
        'DELETE FROM invoice_shot_links WHERE invoice_id=? AND shot_id=?',
        (invoice_id, shot_id)
    )
    conn.commit(); conn.close()
    return jsonify({'status': 'unlinked'})


@bp.route('/api/invoicelog/<int:invoice_id>/shots/bulk', methods=['POST'])
def invoice_shot_bulk(invoice_id):
    """Replace all shot links for an invoice atomically.
    Body: {shot_ids: [int, ...]}
    """
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    shot_ids = (request.json or {}).get('shot_ids', [])
    conn.execute('DELETE FROM invoice_shot_links WHERE invoice_id=?', (invoice_id,))
    for sid in shot_ids:
        try:
            conn.execute(
                'INSERT OR IGNORE INTO invoice_shot_links (invoice_id, shot_id) VALUES (?,?)',
                (invoice_id, int(sid))
            )
        except (ValueError, TypeError):
            pass
    conn.commit(); conn.close()
    return jsonify({'linked': len(shot_ids)})


# ── API – Vendor Performance Score ───────────────────────────────────────────

@bp.route('/api/vendor_performance')
def vendor_performance():
    """Return composite vendor performance scores."""
    conn = core.get_db()
    if not conn: return jsonify([])
    scores = core.get_vendor_perf_scores(conn)
    conn.close()
    return jsonify(scores)


# ── API – All vendors ────────────────────────────────────────────────────────

@bp.route('/api/vendors/all')
def all_vendors():
    conn = core.get_db()
    if not conn: return jsonify([])
    seen = set(); vendors = []
    def add_v(v):
        if v and str(v).strip() and str(v).strip().upper() not in seen:
            seen.add(str(v).strip().upper()); vendors.append(str(v).strip())
    for r in conn.execute('SELECT DISTINCT vendor FROM vendor_tracker WHERE vendor IS NOT NULL').fetchall():
        add_v(r['vendor'])
    for r in conn.execute('SELECT DISTINCT award_vendor FROM assets WHERE award_vendor IS NOT NULL').fetchall():
        for v in str(r['award_vendor']).split(','):
            add_v(v.strip())
    for r in conn.execute('SELECT DISTINCT award FROM bid_compare WHERE award IS NOT NULL').fetchall():
        for v in str(r['award']).split(','):
            add_v(v.strip())
    for r in conn.execute('SELECT DISTINCT vendor FROM vendor_registry WHERE vendor IS NOT NULL').fetchall():
        add_v(r['vendor'])
    conn.close()
    return jsonify(sorted(vendors))


# ── API – Budget Scenario ────────────────────────────────────────────────────

@bp.route('/api/budget_scenario', methods=['GET', 'POST'])
def budget_scenario_collection():
    conn = core.get_db()
    if not conn:
        return (jsonify([]) if request.method == 'GET' else jsonify({'error': 'No project'})), (200 if request.method == 'GET' else 400)
    if request.method == 'GET':
        rows = conn.execute('SELECT * FROM budget_scenario ORDER BY id').fetchall()
        conn.close()
        return jsonify([dict(r) for r in rows])
    d = request.json; c = conn.cursor()
    c.execute('''INSERT INTO budget_scenario (scenario_type, ep, tax_pct, eligible_pct, net_cost, notes, region_mix)
                 VALUES (?,?,?,?,?,?,?)''',
              (d.get('scenario_type','Conservative'), d.get('ep', 101),
               d.get('tax_pct', 0), d.get('eligible_pct', 0),
               d.get('net_cost', 0), d.get('notes', ''), d.get('region_mix', '[]')))
    new_id = c.lastrowid
    conn.commit(); conn.close()
    return jsonify({'id': new_id})


@bp.route('/api/budget_scenario/<int:row_id>', methods=['PUT', 'DELETE'])
def budget_scenario_row(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    if request.method == 'DELETE':
        conn.execute('DELETE FROM budget_scenario WHERE id=?', (row_id,))
        conn.commit(); conn.close()
        return jsonify({'status': 'ok'})
    d = request.json
    allowed = ['scenario_type', 'ep', 'tax_pct', 'eligible_pct', 'net_cost', 'notes', 'region_mix']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE budget_scenario SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [row_id])
        conn.commit()
    conn.close()
    return jsonify({'status': 'ok'})


# ── API – EP Forecast ────────────────────────────────────────────────────────

@bp.route('/api/ep_forecast', methods=['GET'])
def ep_forecast_collection():
    conn = core.get_db()
    if not conn: return jsonify([])
    rows = conn.execute('SELECT * FROM ep_forecast ORDER BY ep').fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@bp.route('/api/ep_forecast/<int:row_id>', methods=['PUT'])
def update_ep_forecast(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    d = request.json
    allowed = ['contingency', 'budget', 'award', 'eligible', 'region_mix', 'global_pct', 'notes']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE ep_forecast SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [row_id])
        conn.commit()
    conn.close()
    return jsonify({'status': 'ok'})


# ── API – Vendor Forecast ────────────────────────────────────────────────────

@bp.route('/api/vendor_forecast', methods=['GET', 'POST'])
def vendor_forecast_collection():
    conn = core.get_db()
    if not conn:
        return (jsonify([]) if request.method == 'GET' else jsonify({'error': 'No project'})), (200 if request.method == 'GET' else 400)
    if request.method == 'GET':
        rows = conn.execute('SELECT * FROM vendor_forecast ORDER BY ep, vendor').fetchall()
        conn.close()
        return jsonify([dict(r) for r in rows])
    d = request.json; c = conn.cursor()
    c.execute('''INSERT INTO vendor_forecast (vendor, ep, gross_local, gross_usd, region, eligible_pct, rebate_pct, notes)
                 VALUES (?,?,?,?,?,?,?,?)''',
              (d.get('vendor',''), d.get('ep', 101), d.get('gross_local', 0), d.get('gross_usd', 0),
               d.get('region',''), d.get('eligible_pct', 0), d.get('rebate_pct', 0), d.get('notes','')))
    new_id = c.lastrowid
    conn.commit(); conn.close()
    return jsonify({'id': new_id})


@bp.route('/api/vendor_forecast/<int:row_id>', methods=['PUT', 'DELETE'])
def vendor_forecast_row(row_id):
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400
    if request.method == 'DELETE':
        conn.execute('DELETE FROM vendor_forecast WHERE id=?', (row_id,))
        conn.commit(); conn.close()
        return jsonify({'status': 'ok'})
    d = request.json
    allowed = ['vendor', 'ep', 'gross_local', 'gross_usd', 'region', 'eligible_pct', 'rebate_pct', 'notes']
    updates = {k: d[k] for k in allowed if k in d}
    if updates:
        sql = 'UPDATE vendor_forecast SET ' + ', '.join(f'{k}=?' for k in updates) + ' WHERE id=?'
        conn.execute(sql, list(updates.values()) + [row_id])
        conn.commit()
    conn.close()
    return jsonify({'status': 'ok'})


# ── API – Data Graph ─────────────────────────────────────────────────────────

@bp.route('/api/data_graph', methods=['GET'])
def get_data_graph():
    """Return enriched shot data with cross-module links."""
    conn = core.get_db()
    if not conn: return jsonify([])
    from services.data_graph import build_shot_graph
    ep = request.args.get('ep', type=int)
    try:
        data = build_shot_graph(conn, ep)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@bp.route('/api/season_risk_context', methods=['GET'])
def get_season_risk_context():
    """Return season-level risk context for LLM prompts."""
    conn = core.get_db()
    if not conn: return jsonify({})
    from services.data_graph import get_season_risk_context
    try:
        return jsonify(get_season_risk_context(conn))
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@bp.route('/api/llm/scenario_advice', methods=['POST'])
def scenario_advice():
    """
    AI recommendation for budget scenario selection for an episode.
    POST body: { ep }
    """
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400

    import llm as llm_module
    from services.llm_helpers import complete_json
    from services.data_graph import get_season_risk_context

    d  = request.json or {}
    ep = d.get('ep')
    if not ep:
        return jsonify({'error': 'ep required'}), 400

    scenarios = conn.execute(
        'SELECT * FROM budget_scenario WHERE ep=? ORDER BY id', (ep,)
    ).fetchall()
    ep_agg = conn.execute(
        "SELECT SUM(COALESCE(cost_est,0)) as est, SUM(COALESCE(efc,0)) as efc "
        "FROM shots WHERE ep=? AND omit=0", (ep,)
    ).fetchone()

    current_est = ep_agg['est'] or 0
    current_efc = ep_agg['efc'] or 0
    variance    = current_efc - current_est
    var_pct     = f"{(variance / current_est * 100):+.1f}" if current_est else "N/A"

    scenarios_text = '\n'.join(
        f"  [{r['scenario_type']}] tax={r['tax_pct']}% eligible={r['eligible_pct']}% "
        f"net_cost=${r['net_cost']:,.0f} notes={r['notes'] or ''}"
        for r in scenarios
    ) or '  (no scenarios defined)'

    risk_ctx = get_season_risk_context(conn)
    conn.close()

    risk_summary = f"Season variance: ${risk_ctx['variance']:,.0f} across {risk_ctx['total_shots']} shots"

    ctx = {
        'ep':           ep,
        'current_efc':  f"{current_efc:,.0f}",
        'current_est':  f"{current_est:,.0f}",
        'variance':     f"{variance:+,.0f}",
        'variance_pct': var_pct,
        'scenarios_text': scenarios_text,
        'risk_context': risk_summary,
    }
    prompt = llm_module.build_prompt('scenario_advice', ctx)
    if not prompt:
        return jsonify({'error': 'prompt build failed'}), 400

    client = core.get_llm_client()
    result = complete_json(prompt, client,
                           cache_task='scenario_advice',
                           cache_ctx={'ep': str(ep)})
    if result is None:
        return jsonify({'error': 'LLM unavailable'}), 503
    return jsonify(result)


@bp.route('/api/llm/vendor_narrative', methods=['POST'])
def vendor_narrative():
    """
    AI narrative summary of vendor health for the season. Streams response.
    """
    conn = core.get_db()
    if not conn: return jsonify({'error': 'No project'}), 400

    import llm as llm_module
    from services.llm_helpers import stream_llm
    from services.vendors_svc import get_vendor_loads, get_vendor_perf_scores
    from services.data_graph import get_season_risk_context
    from flask import session

    loads       = get_vendor_loads(conn)
    scores_list = get_vendor_perf_scores(conn)
    scores      = {r['vendor']: r for r in scores_list}
    risk_ctx    = get_season_risk_context(conn)

    # Build vendor list text
    top_v = conn.execute(
        "SELECT award_vendor as vendor, COUNT(*) as shots, "
        "SUM(COALESCE(efc,0)) as efc FROM shots "
        "WHERE omit=0 AND award_vendor IS NOT NULL AND award_vendor!='' "
        "GROUP BY award_vendor ORDER BY efc DESC LIMIT 8"
    ).fetchall()
    vendor_lines = []
    flags_lines  = []
    for v in top_v:
        vname = v['vendor']
        load  = loads.get(vname.upper(), {})
        score = scores.get(vname.upper(), {})
        util  = load.get('utilization')
        perf  = score.get('composite_score')
        vendor_lines.append(
            f"  {vname}: {v['shots']} shots EFC=${v['efc']:,.0f}"
            f"{f' util={util*100:.0f}%' if util is not None else ''}"
            f"{f' perf={perf:.2f}' if perf is not None else ''}"
        )
        if util and util > 0.85:
            flags_lines.append(f"  CAPACITY: {vname} at {util*100:.0f}%")
        if perf is not None and perf < 0.4:
            flags_lines.append(f"  LOW PERF: {vname} score {perf:.2f}")

    conn.close()

    ctx = {
        'project':       session.get('project_name', 'Unknown'),
        'season':        session.get('project_season', ''),
        'vendor_list':   '\n'.join(vendor_lines) or '  (no vendor data)',
        'total_shots':   risk_ctx['total_shots'],
        'season_efc':    f"{risk_ctx['total_efc']:,.0f}",
        'season_est':    f"{risk_ctx['total_est']:,.0f}",
        'variance':      f"{risk_ctx['variance']:+,.0f}",
        'flags_text':    '\n'.join(flags_lines) or '  (none)',
    }
    prompt = llm_module.build_prompt('vendor_narrative', ctx)
    if not prompt:
        return jsonify({'error': 'prompt build failed'}), 400

    client = core.get_llm_client()
    return stream_llm(prompt, client,
                      cache_task='vendor_narrative',
                      cache_ctx={'season': session.get('project_season', '')})
